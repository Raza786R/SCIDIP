"""Create the compact Phase 2 DOCX teaching pack from reviewed Markdown sources."""
import json,re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
root=Path(__file__).resolve().parents[2];sources=root/'phase2/sources';out=root/'phase2/docs';out.mkdir(exist_ok=True)
refs=json.loads((sources/'references.json').read_text())
aero=root/'phase2/results/aerospace_v2';sona=root/'phase2/results/sona_v2'
a=json.loads((aero/'best_candidate.json').read_text());s=json.loads((sona/'best_candidate.json').read_text());met=json.loads((aero/'metrics.json').read_text())

def dynamic_table(name):
 if name=='current_results':
  return [['Screening quantity','Aerospace hypothesis','Sona coupon hypothesis'],
   ['Candidate ID',a['composition_id'],s['composition_id']],
   ['Peak in plane stress MPa',f"{a['screening']['peak_inplane_stress_MPa']:.1f}",f"{s['screening']['peak_inplane_stress_MPa']:.1f}"],
   ['Density estimate g per cubic cm',f"{a['mixture_estimates']['density_g_cm3']:.3f}",f"{s['mixture_estimates']['density_g_cm3']:.3f}"],
   ['Thickness mm',f"{a['design']['thickness_m']*1000:.3f}",f"{s['design']['thickness_m']*1000:.3f}"],
   ['Compression YS prediction MPa',f"{a['screening']['compression_YS_MPa_prediction']:.1f}",f"{s['screening']['compression_YS_MPa_prediction']:.1f}"],
   ['Eligible strength lower bound MPa','0.0','0.0'],['Qualification','Unmeasured screen','Unmeasured screen']]
 if name=='model_metrics':
  rows=[['Target','Train rows and groups','Calibration rows and groups','Test rows and groups','Test MAE and R squared']]
  for key,label in [('bulk','Bulk modulus GPa'),('yield','Compression YS MPa')]:
   m=met[key];rows.append([label,*[f"{m['rows_by_partition'][p]} rows / {m['groups_by_partition'][p]} groups" for p in ['train','calibration','test']],f"{m['locked_test']['MAE']:.3f} / {m['locked_test']['R2']:.3f}"])
  return rows
 if name=='control_results':
  frame=pd.read_csv(root/'phase2/results/graded_bilayer_controls.csv')
  return [['Profile','Constraint','Architecture','Peak stress MPa','Heat flux MW per square m']]+[[r.profile.replace('_v2','').replace('aerospace','Aerospace').replace('sona','Sona'),r.constraint.replace('_',' '),r.architecture.replace('_',' '),f'{r.peak_inplane_stress_MPa:.1f}',f'{r.heat_flux_W_m2/1e6:.3f}'] for r in frame.itertuples()]
 if name=='composition_recipe':
  return [['Element','Atomic percent','Metal mass percent']]+[[el,f'{100*a["atomic_fractions"][el]:.4f}',f'{100*a["mixture_estimates"]["mass_fractions"][el]:.4f}'] for el in ['Mo','Nb','Ta','W','Cr','V','Ti','Zr','Hf']]
 if name=='layer_recipe':
  f=pd.read_csv(aero/'coupon_recipe_11_layers.csv')
  return [['Layer','Depth start to end mm','Ceramic volume percent','Metal mass percent']]+[[str(r.layer),f'{r.z_start_m*1000:.3f} to {r.z_end_m*1000:.3f}',f'{r.ceramic_volume_fraction*100:.2f}',f'{r.metal_mass_fraction*100:.2f}'] for r in f.itertuples()]
 if name=='budget_heads':
  return [['Head','INR lakh','Percent of total'],*[[h,f'{v:.2f}',f'{v/65*100:.2f}'] for h,v in [('Equipment',8),('Staffing',17),('Consumables',14),('Outsourcing',19),('Patent support',3),('Contingency',4)]],['Total','65.00','100.00']]
 if name=='milestones':
  return [['Gate and months','Evidence to deliver','INR lakh'],['1 Months 1 to 4','Agreed use case, benchmark, process and thermodynamic feasibility, test protocol','10.80'],['2 Months 5 to 10','Repeatable coupons and actual chemistry phase and defect records, or justified pivot','19.10'],['3 Months 11 to 18','Measured thermal environmental and relevant mechanical comparison with controls','21.30'],['4 Months 19 to 24','Prospective feedback iteration, application and economic go or no go, data handover','13.80'],['Total','Planned expenditure, not approved release tranches','65.00']]
 raise ValueError(name)


def inline(p,text):
 # Minimal Markdown emphasis and code; citations remain ordinary source labels.
 for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text):
  if part.startswith('**') and part.endswith('**'):p.add_run(part[2:-2]).bold=True
  elif part.startswith('`') and part.endswith('`'):
   r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(10)
  else:p.add_run(part)


def table(doc,rows):
 n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 width=7.0
 if n==2: widths=[1.8,5.2]
 elif n==3:widths=[2.45,2.275,2.275]
 elif n==4:widths=[.7,2.7,1.8,1.8] if rows[0][0]=='Layer' else [1.25,2.0,1.6,2.15]
 elif n==5:widths=[1.3,1.45,1.45,1.4,1.4]
 else:widths=[width/n]*n
 if n==3 and rows[0][0]=='Gate and months':widths=[1.3,4.7,1.0]
 if n==3 and rows[0][0]=='Head':widths=[3.6,1.5,1.9]
 if n==3 and rows[0][0]=='Element':widths=[1.1,2.95,2.95]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for ri,row in enumerate(rows):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  for j,text in enumerate(row):
   c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcPr.append(b);m=OxmlElement('w:tcMar')
   for k in ('top','bottom','left','right'):
    e=OxmlElement('w:'+k);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
   tcPr.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'23445B' if ri==0 else ('F1F5F8' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.04
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or (len(text)<20 and re.search(r'\d',text) and j>0) else WD_ALIGN_PARAGRAPH.LEFT
   r=p.add_run(text);r.font.name='Calibri';r.font.size=Pt(10.3);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
  trPr=t.rows[ri]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=.5;p.add_run().font.size=Pt(2)


def image(doc,name,caption,equation=False):
 path=root/'phase2/assets'/name
 with Image.open(path) as im:w,h=im.size
 width=6.8 if not equation else min(6.8,max(3.6,w/330))
 if not equation and h/w*width>3.9:width=3.9*w/h
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 shape=p.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',caption)
 p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.space_after=Pt(9)


manifest=[]
for source in sorted(sources.glob('[0-9][0-9]_*.md')):
 content=source.read_text();doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.left_margin=sec.right_margin=Inches(.75);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.header_distance=sec.footer_distance=Inches(.25)
 styles=doc.styles
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','Header','Footer']:
  st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 styles['Normal'].font.size=Pt(11);styles['Normal'].paragraph_format.line_spacing=1.08;styles['Normal'].paragraph_format.space_after=Pt(6)
 styles['Title'].font.size=Pt(22);styles['Title'].font.bold=True;styles['Title'].paragraph_format.space_after=Pt(8)
 styles['Subtitle'].font.size=Pt(10.5);styles['Subtitle'].paragraph_format.space_after=Pt(12)
 styles['Heading 1'].font.size=Pt(14);styles['Heading 1'].font.bold=True;styles['Heading 1'].paragraph_format.space_before=Pt(12);styles['Heading 1'].paragraph_format.space_after=Pt(6);styles['Heading 1'].paragraph_format.keep_with_next=True
 styles['Heading 2'].font.size=Pt(12);styles['Heading 2'].font.bold=True
 styles['Caption'].font.size=Pt(9.5);styles['Caption'].font.italic=True
 for st in styles:
  if st.type==1:
   st.paragraph_format.widow_control=True
   for b in list(st.element.xpath('./w:pPr/w:pBdr')):b.getparent().remove(b)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run('Crystal Paradigm   |   ');r.font.size=Pt(8)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
 lines=content.splitlines();i=0;first=True
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('# '):doc.add_paragraph(line[2:],'Title');first=True
  elif first:doc.add_paragraph(line.replace('|',' '),'Subtitle');first=False
  elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
  elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
  elif line.startswith('|'):
   rows=[[c.strip() for c in line.strip('|').split('|')]]
   while i<len(lines) and lines[i].strip().startswith('|'):
    row=[c.strip() for c in lines[i].strip().strip('|').split('|')];i+=1
    if not all(re.fullmatch(r'[-: ]+',c) for c in row):rows.append(row)
   table(doc,rows)
  elif line.startswith('{{table:'):table(doc,dynamic_table(line[8:-2]))
  elif line.startswith('{{figure:') or line.startswith('{{equation:'):
   kind,rest=line[2:-2].split(':',1);name,caption=rest.split('|',1);image(doc,name,caption,kind=='equation')
  else:
   p=doc.add_paragraph(style='Normal');inline(p,line)
 used=sorted(set(re.findall(r'\bS\d+\b',content)),key=lambda x:int(x[1:]))
 if used:
  doc.add_paragraph('Sources and evidence','Heading 1')
  for key in used:
   r=refs[key];p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
   x=p.add_run(f"[{key}] {r['title']}. ");x.bold=True;x.font.size=Pt(9.5)
   x=p.add_run(r['url']);x.font.size=Pt(9.5)
 doc.core_properties.title=lines[0][2:];doc.core_properties.author='Crystal Paradigm';doc.core_properties.subject='HE MCGA Phase 2 technical learning and SCIDIP preparation'
 dest=out/(source.stem+'.docx');doc.save(dest)
 manifest.append({'file':dest.name,'source':str(source.relative_to(root)),'word_count':len(content.split())})
(root/'phase2/document_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(manifest,indent=2))
