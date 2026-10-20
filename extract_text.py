import os
import docx
import pptx

def extract_docx(filepath, outpath):
    doc = docx.Document(filepath)
    text = []
    for p in doc.paragraphs:
        if p.text.strip():
            text.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                if cell.text.strip():
                    row_data.append(cell.text.strip().replace('\n', ' '))
            if row_data:
                text.append(" | ".join(row_data))
    with open(outpath, 'w', encoding='utf-8') as f:
        f.write('\n'.join(text))

def extract_pptx(filepath, outpath):
    prs = pptx.Presentation(filepath)
    text = []
    for i, slide in enumerate(prs.slides):
        text.append(f"\n--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text)
    with open(outpath, 'w', encoding='utf-8') as f:
        f.write('\n'.join(text))

base_dir = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\references_and_guides\FINAL_SUBMISSION_FILES"
extract_docx(os.path.join(base_dir, "Computational_Document_Analysis_REVISED.docx"), os.path.join(base_dir, "docx_text.txt"))
extract_pptx(os.path.join(base_dir, "CRYSTAL PARADIGM_ FOR HOTTER, LONGER-LIVED PROPULSION MATERIALS.pptx"), os.path.join(base_dir, "pptx_text.txt"))
print("Extraction complete.")
