import os
from docx import Document

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
DOCX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_TechnicalAppendix.docx")

doc = Document(DOCX_PATH)

# 1. Update Indrasena's Profile
for p in doc.paragraphs:
    if "K. Indrasena Reddy is a Computer Science Engineering" in p.text:
        # Rebuild his profile with HE-MCGA
        full_text = "K. Indrasena Reddy is a Computer Science Engineering student at Neil Gogte Institute of Technology (NGIT). He previously served as an AI Developer Intern at Viswam.AI, where he built natural language processing applications and contributed to experimental language model research. Currently, he advances materials informatics and graph neural network modeling for large-scale property prediction as an AI Research Intern at Crystal Paradigm, where he architected the materials informatics and validation pipelines for the HE-MCGA aerospace project."
        
        p.clear()
        
        highlights = [
            "AI Developer Intern at Viswam.AI",
            "natural language processing applications",
            "architected the materials informatics and validation pipelines for the HE-MCGA"
        ]
        
        current_text = full_text
        for h in highlights:
            if h in current_text:
                parts = current_text.split(h)
                p.add_run(parts[0])
                r = p.add_run(h)
                r.bold = True
                if len(parts) > 1:
                    current_text = parts[1]
        
        if current_text:
            p.add_run(current_text)

        p.paragraph_format.keep_together = True

# 2. Remove the 18 interns block
for p in doc.paragraphs:
    if "EXTENDED AI EXECUTION TEAM" in p.text or "18 undergraduate AI Research Interns" in p.text:
        p.clear() # clear text
        p._element.getparent().remove(p._element) # remove empty paragraph

doc.save(DOCX_PATH)
print("Updated Indrasena's profile and removed 18 interns.")
