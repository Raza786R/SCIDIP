"""Extract exact style data from Final.pptx slides 1-14 so we can perfectly match them."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

p = Presentation("CrystalParadigm_SCIDIP_PitchDeck_Final.pptx")

for idx, sl in enumerate(p.slides):
    print(f"\n{'='*80}")
    print(f"SLIDE {idx+1} — shapes: {len(sl.shapes)}")
    print(f"{'='*80}")
    
    # background
    try:
        bg = sl.background.fill
        if bg.type is not None:
            print(f"  BG fill type: {bg.type}")
            try: print(f"  BG color: #{bg.fore_color.rgb}")
            except: pass
    except: pass
    
    for s in sl.shapes:
        l = round(s.left/914400, 2) if s.left else 0
        t = round(s.top/914400, 2) if s.top else 0
        w = round(s.width/914400, 2) if s.width else 0
        h = round(s.height/914400, 2) if s.height else 0
        
        # fill info
        fill_info = ""
        try:
            sf = s.fill
            if sf.type is not None:
                try: fill_info = f" fill=#{sf.fore_color.rgb}"
                except: fill_info = f" fill=type{sf.type}"
        except: pass
        
        # line info
        line_info = ""
        try:
            if s.line.color.rgb:
                line_info = f" border=#{s.line.color.rgb}"
        except: pass
        
        print(f"  [{s.shape_type}] '{s.name}' l={l} t={t} w={w} h={h}{fill_info}{line_info}")
        
        if hasattr(s, 'text_frame'):
            tf = s.text_frame
            for pi, para in enumerate(tf.paragraphs[:8]):
                txt = para.text.strip()
                if not txt: continue
                align = str(para.alignment).split('.')[-1] if para.alignment else 'LEFT'
                sz_list = []
                col_list = []
                bold_list = []
                for run in para.runs:
                    if run.font.size: sz_list.append(round(run.font.size/12700))
                    try:
                        if run.font.color.rgb: col_list.append(f"#{run.font.color.rgb}")
                    except: pass
                    if run.font.bold: bold_list.append('B')
                sz = sz_list[0] if sz_list else '?'
                col = col_list[0] if col_list else '?'
                bold = 'BOLD' if bold_list else ''
                print(f"    p{pi}: sz={sz} col={col} {bold} align={align} | '{txt[:60]}'")
