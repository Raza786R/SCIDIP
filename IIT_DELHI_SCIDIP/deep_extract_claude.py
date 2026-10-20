"""
Deep extractor for the 'by claude' PPTX - extracts exact shape positions,
sizes, colors, text, and types for every slide so we can faithfully
reproduce the best visualizations in the new deck.
"""
import sys, codecs, json
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import traceback

fname = 'CrystalParadigm_SCIDIP_PitchDeck(by claude).pptx'
prs = Presentation(fname)

slide_data = []
for i, slide in enumerate(prs.slides):
    shapes = []
    for shape in slide.shapes:
        s = {
            'name': shape.name,
            'shape_type': str(shape.shape_type),
            'left': round(shape.left / 914400, 3),   # convert EMU to inches
            'top': round(shape.top / 914400, 3),
            'width': round(shape.width / 914400, 3),
            'height': round(shape.height / 914400, 3),
        }
        if hasattr(shape, 'text') and shape.text.strip():
            s['text'] = shape.text.strip()[:300]
        
        # Get fill color if available
        try:
            fill = shape.fill
            if fill.type is not None:
                if hasattr(fill.fore_color, 'rgb'):
                    s['fill_color'] = str(fill.fore_color.rgb)
        except:
            pass
        
        # Get font info from first paragraph if text frame
        try:
            if shape.has_text_frame:
                para = shape.text_frame.paragraphs[0]
                if para.runs:
                    run = para.runs[0]
                    if run.font.size:
                        s['font_size'] = round(run.font.size / 12700)
                    if run.font.bold:
                        s['font_bold'] = True
                    try:
                        if run.font.color.rgb:
                            s['font_color'] = str(run.font.color.rgb)
                    except:
                        pass
        except:
            pass

        shapes.append(s)
    
    slide_data.append({
        'slide_num': i+1,
        'shape_count': len(shapes),
        'shapes': shapes
    })

with open('by_claude_deep_analysis.json', 'w', encoding='utf-8') as f:
    json.dump(slide_data, f, indent=2, ensure_ascii=False)

print(f"Done! Extracted {len(slide_data)} slides to by_claude_deep_analysis.json")
for s in slide_data:
    print(f"Slide {s['slide_num']}: {s['shape_count']} shapes")
