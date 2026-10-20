# Final extraction of modified PPTX for review
from pptx import Presentation
import os

PPTX_PATH = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP\CrystalParadigm_SCIDIP_PitchDeck.pptx"

prs = Presentation(PPTX_PATH)
output = []

for slide_num, slide in enumerate(prs.slides, 1):
    output.append(f"\n{'='*60}")
    output.append(f"SLIDE {slide_num}")
    output.append('='*60)
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    output.append(text)
        if shape.has_table:
            output.append("[TABLE]")
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                output.append(" | ".join(cells))

print("\n".join(output))
