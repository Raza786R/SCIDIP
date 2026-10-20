import os
from docx import Document

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
DOCX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_TechnicalAppendix.docx")

doc = Document(DOCX_PATH)

sec11_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "11." in p.text and "Ask" in p.text:
        sec11_idx = i
        
        # 1. Remove Page Break Before property
        p.paragraph_format.page_break_before = False
        p.paragraph_format.space_before = 0
        
        # 2. Remove manual hard page breaks inside the title
        for run in p.runs:
            if "page" in run._element.xml and "br" in run._element.xml:
                run.clear()
        break

if sec11_idx != -1:
    # 3. Delete all empty paragraphs immediately preceding Section 11
    for i in range(sec11_idx - 1, -1, -1):
        p = doc.paragraphs[i]
        if not p.text.strip():
            p.clear()
            p._element.getparent().remove(p._element)
        else:
            # Add exactly one blank line before Section 11 for breathing room
            p.insert_paragraph_after("")
            break

doc.save(DOCX_PATH)
print("Gap before Section 11 fixed.")
