import sys, codecs
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
from pptx import Presentation

files = {
    'FINAL': 'CrystalParadigm_SCIDIP_PitchDeck_Final.pptx',
    'BY_CLAUDE': 'CrystalParadigm_SCIDIP_PitchDeck(by claude).pptx',
    'V2': 'CrystalParadigm_SCIDIP_PitchDeck_v2.pptx'
}

for label, fname in files.items():
    try:
        prs = Presentation(fname)
        print(f'=== {label}: {fname} ===')
        print(f'Slides: {len(prs.slides)}')
        for i, slide in enumerate(prs.slides):
            texts = []
            shapes_info = []
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    t = shape.text.strip().replace('\n',' ')[:120]
                    texts.append(t)
                shapes_info.append(f'{shape.shape_type}')
            print(f'  Slide {i+1} ({len(slide.shapes)} shapes): {" | ".join(texts[:5])}')
        print()
    except Exception as e:
        print(f'Error {label}: {e}')
