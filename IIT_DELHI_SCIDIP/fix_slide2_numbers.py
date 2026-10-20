# -*- coding: utf-8 -*-
from pptx import Presentation
import os

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
PPTX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_PitchDeck.pptx")

prs = Presentation(PPTX_PATH)
fixes = 0

def process_shape(shape):
    global fixes
    if getattr(shape, 'shape_type', None) == 6:  # Group
        for s in shape.shapes:
            process_shape(s)
    
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            original = p.text
            new_text = original
            
            # SLIDE 2 FIXES
            new_text = new_text.replace("CHALLENGE #1", "FOCUS AREA 1")
            new_text = new_text.replace("SMART MATERIALS", "ADVANCED MATERIALS")
            new_text = new_text.replace("[1.1.1 & 1.2.2] ", "")
            new_text = new_text.replace("[1.1.2] ", "")
            new_text = new_text.replace("[1.3.2] ", "")
            
            new_text = new_text.replace("CHALLENGE #2", "FOCUS AREA 2")
            new_text = new_text.replace("INTELLIGENCE FOR MATERIALS & PROCESSES", "INDUSTRY 4.0 & AI")
            new_text = new_text.replace("AI-Driven Design, Simulation & Chemical Gradients", "AI-Driven Design, Simulation & Digital Twins")
            new_text = new_text.replace("[2.3.1] ", "")
            new_text = new_text.replace("[2.2.1] ", "")
            new_text = new_text.replace("[2.1.1 & 2.1.2] ", "")
            
            new_text = new_text.replace("(Ch 1)", "(Focus Area 1)")
            new_text = new_text.replace("(Ch 2)", "(Focus Area 2)")

            if new_text != original:
                if p.runs:
                    p.runs[0].text = new_text
                    for r in p.runs[1:]:
                        r.text = ""
                    fixes += 1
                    print(f"Fixed: {original[:50]} -> {new_text[:50]}")

for slide in prs.slides:
    for shape in slide.shapes:
        process_shape(shape)

prs.save(PPTX_PATH)
print(f"Total precise PPTX fixes applied: {fixes}")
