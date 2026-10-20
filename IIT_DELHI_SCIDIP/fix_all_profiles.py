import os
from docx import Document
from docx.shared import Pt, RGBColor

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
DOCX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_TechnicalAppendix.docx")

doc = Document(DOCX_PATH)

# Team member texts
team_members = [
    {
        "title": "AI/ML Lead: Dr. Shilpa Choudhary",
        "desc": "Dr. Choudhary, recognized among the Top 2% Scientists Worldwide by Stanford University (2024), brings over 16 years of research expertise in Artificial Intelligence, Deep Learning, and AI-driven scientific applications. With postdoctoral research experience at the Singapore Institute of Technology and the National Kaohsiung University of Science and Technology (Taiwan), she has authored over 100 research papers (1763 citations, H-index 22), 12 academic books, and holds multiple patents in emerging technology domains. At Crystal Paradigm, she leads AI innovation, ML Architecture, Constrained Optimization, and Surrogate Model Design.",
        "highlights": [
            "Top 2% Scientists Worldwide by Stanford University (2024)",
            "16 years of research expertise in Artificial Intelligence, Deep Learning, and AI-driven scientific applications",
            "postdoctoral research experience at the Singapore Institute of Technology and the National Kaohsiung University of Science and Technology (Taiwan)",
            "100 research papers",
            "12 academic books",
            "multiple patents"
        ]
    },
    {
        "title": "Technical Lead: Dr. Kasarapu Venkataramana",
        "desc": "Dr. Venkataramana, an Associate Fellow of the Telangana Academy of Sciences (2024), brings 12 years of research and 5 years of academic teaching experience to the team. He has authored 38 research publications (610+ citations, h-index 14) and 2 book chapters, while serving as a recognized reviewer for prestigious scientific journals across Elsevier, Springer, and Taylor & Francis. In his current role at CSIR-IICT, he is leading DRDO-sponsored pyrotechnic nanomaterial fabrication projects. He provides critical domain oversight on computational chemistry, thermodynamics, and physical experimental validation.",
        "highlights": [
            "Associate Fellow of the Telangana Academy of Sciences (2024)",
            "12 years of research",
            "38 research publications",
            "2 book chapters",
            "leading DRDO-sponsored pyrotechnic nanomaterial fabrication projects"
        ]
    },
    {
        "title": "TEAM MEMBER 1: MOHAMMED AHMED RAZA",
        "desc": "Mohammed Ahmed Raza is a Computer Science and Engineering student at Neil Gogte Institute of Technology (NGIT) and a Certified AI/ML Professional from iHub-Data, IIIT Hyderabad. He previously served as a Winter Research Intern at IIIT Hyderabad, where he engineered advanced computer vision systems for UAV spatial awareness and obstacle detection. Currently, he advances materials science and deep learning innovations as an AI Research Intern at Crystal Paradigm, developing inverse design frameworks for HE-MCGA, computationally engineering next-generation superalloys for extreme aerospace applications. Most recently, he co-authored research introducing a breakthrough Attention Graph Siamese Transformer (DPT-GNN) that drastically optimizes lithium-ion battery capacity predictions through transition-aware learning.",
        "highlights": [
            "Certified AI/ML Professional from iHub-Data, IIIT Hyderabad",
            "Winter Research Intern at IIIT Hyderabad",
            "advanced computer vision systems",
            "developing inverse design frameworks for HE-MCGA",
            "Attention Graph Siamese Transformer (DPT-GNN)"
        ]
    },
    {
        "title": "TEAM MEMBER 2: K. INDRASENA REDDY",
        "desc": "K. Indrasena Reddy is a Computer Science Engineering student at Neil Gogte Institute of Technology (NGIT). He previously served as an AI Developer Intern at Viswam.AI, where he built natural language processing applications and contributed to experimental language model research. Currently, he advances materials informatics and graph neural network modeling for large-scale property prediction as an AI Research Intern at Crystal Paradigm, where he architected the materials informatics and validation pipelines for the HE-MCGA aerospace project.",
        "highlights": [
            "AI Developer Intern at Viswam.AI",
            "natural language processing applications",
            "architected the materials informatics and validation pipelines for the HE-MCGA"
        ]
    }
]

# 1. Clean up "CRYSTALPARADIGM PRIVATE LIMITED" spacing
for i, p in enumerate(doc.paragraphs):
    if "CRYSTALPARADIGM PRIVATE LIMITED" in p.text.upper():
        p.insert_paragraph_before("")
        break

# 2. Delete ALL garbage profiles between "Email:" and "For inquiries"
start_del = -1
end_del = -1

for i, p in enumerate(doc.paragraphs):
    if "khub.crystalparadigm@gmail.com" in p.text:
        start_del = i + 1
    if "For inquiries regarding collaboration" in p.text:
        end_del = i

if start_del != -1 and end_del != -1 and end_del > start_del:
    # Delete backwards to not mess up indexes
    for i in range(end_del - 1, start_del - 1, -1):
        p = doc.paragraphs[i]
        p.clear()
        p._element.getparent().remove(p._element)

# 3. Re-insert perfectly formatted
insert_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "For inquiries regarding collaboration" in p.text:
        insert_idx = i
        break

if insert_idx != -1:
    p_ref = doc.paragraphs[insert_idx]
    
    # Add spacing after Email
    p_ref.insert_paragraph_before("")
    
    for member in team_members:
        p_title = p_ref.insert_paragraph_before("")
        run = p_title.add_run(member["title"])
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 51, 102)
        
        p_ref.insert_paragraph_before("")
        
        p_desc = p_ref.insert_paragraph_before("")
        current_text = member["desc"]
        
        for h in member["highlights"]:
            if h in current_text:
                parts = current_text.split(h)
                p_desc.add_run(parts[0])
                r = p_desc.add_run(h)
                r.bold = True
                if len(parts) > 1:
                    current_text = parts[1]
        
        if current_text:
            p_desc.add_run(current_text)
            
        p_desc.paragraph_format.keep_together = True
        p_title.paragraph_format.keep_with_next = True
        
        p_ref.insert_paragraph_before("")
        p_ref.insert_paragraph_before("")

    # 4. Insert the "and team" smart block safely without numbers
    p_team_title = p_ref.insert_paragraph_before("")
    run_t = p_team_title.add_run("EXTENDED RESEARCH TEAM")
    run_t.bold = True
    run_t.font.size = Pt(12)
    run_t.font.color.rgb = RGBColor(0, 51, 102)
    p_team_title.paragraph_format.keep_with_next = True
    
    p_ref.insert_paragraph_before("")
    
    p_team_desc = p_ref.insert_paragraph_before("The core leadership is supported by a dedicated extended cohort of AI Research Interns executing high-throughput computational modeling, data processing, and multi-physics validations across our platform workflows.")
    p_team_desc.paragraph_format.keep_together = True
    
    p_ref.insert_paragraph_before("")
    p_ref.insert_paragraph_before("")

doc.save(DOCX_PATH)
print("Profiles perfectly rebuilt without hallucinations, and safe team block added.")
