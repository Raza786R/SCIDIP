"""Quick audit of the upgraded IIT Delhi PPTX."""
from pptx import Presentation
import sys, codecs
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')

prs = Presentation("CrystalParadigm_SCIDIP_IITDelhi_FINAL.pptx")
print(f"Slides: {len(prs.slides)}  |  Size: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
print(f"{'='*95}")

for i, slide in enumerate(prs.slides):
    shapes = list(slide.shapes)
    texts = [s.text.strip() for s in shapes if hasattr(s, 'text') and s.text.strip()]
    total_chars = sum(len(t) for t in texts)
    
    # Find first meaningful title text
    title_text = ""
    for s in shapes:
        if hasattr(s, 'text') and s.text.strip() and len(s.text.strip()) > 8:
            title_text = s.text.strip()[:75].replace('\n',' | ')
            break
    
    is_new = "NEW" if i >= 13 and i < 19 else "   "
    print(f"Slide {i+1:2d} [{is_new}]  Shapes: {len(shapes):3d}  Chars: {total_chars:5d}  '{title_text}'")

print(f"{'='*95}")
