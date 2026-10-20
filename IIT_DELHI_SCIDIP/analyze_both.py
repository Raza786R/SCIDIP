"""
Deep analysis of BOTH reference PPTX files:
1. CrystalParadigm_SCIDIP_PitchDeck(by claude).pptx - for ELEMENTS/VISUALIZATIONS
2. CrystalParadigm_SCIDIP_PitchDeck_Final.pptx - for THEME/CONTENT

Outputs a full slide-by-slide breakdown to understand exactly what's in each.
"""
import sys, codecs, json
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

def analyze_pptx(fname, label):
    prs = Presentation(fname)
    print(f"\n{'='*70}")
    print(f"FILE: {label} ({fname})")
    print(f"Slides: {len(prs.slides)}, Size: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
    
    # Get background color of first slide
    try:
        bg = prs.slides[0].background.fill
        print(f"Slide BG type: {bg.type}")
        if hasattr(bg, 'fore_color'):
            try:
                print(f"Slide BG color: #{bg.fore_color.rgb}")
            except: pass
    except: pass
    
    print(f"{'='*70}")
    
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        
        # Background
        try:
            sbg = slide.background.fill
            if hasattr(sbg, 'fore_color'):
                try:
                    print(f"  BG: #{sbg.fore_color.rgb}")
                except: print(f"  BG: type={sbg.type}")
        except: pass
        
        shapes_data = []
        for shape in slide.shapes:
            s = {
                'name': shape.name,
                'type': str(shape.shape_type),
                'l': round(shape.left/914400, 2),
                't': round(shape.top/914400, 2),
                'w': round(shape.width/914400, 2),
                'h': round(shape.height/914400, 2),
            }
            
            # Fill
            try:
                if shape.fill.type is not None:
                    try:
                        s['fill'] = f"#{shape.fill.fore_color.rgb}"
                    except: s['fill'] = f"type={shape.fill.type}"
            except: pass
            
            # Text
            if hasattr(shape, 'text') and shape.text.strip():
                s['text'] = shape.text.strip()[:80].replace('\n', ' | ')
                # Font info
                try:
                    para = shape.text_frame.paragraphs[0]
                    if para.runs:
                        run = para.runs[0]
                        if run.font.size:
                            s['fs'] = round(run.font.size/12700)
                        try:
                            s['fc'] = f"#{run.font.color.rgb}"
                        except: pass
                        if run.font.bold:
                            s['bold'] = True
                except: pass
            
            # Image
            if 'PICTURE' in str(shape.shape_type):
                s['type'] = 'IMAGE'
            
            shapes_data.append(s)
        
        # Print shapes
        for s in shapes_data:
            txt = s.get('text', '')
            fill = s.get('fill', '')
            fc = s.get('fc', '')
            fs = s.get('fs', '')
            info = f"[{s['w']:.1f}x{s['h']:.1f} @ ({s['l']:.1f},{s['t']:.1f})]"
            print(f"  {s['type'][:8]:8} {info:25} fill={fill:10} fc={fc:10} fs={fs:3}  '{txt}'")

# Analyze both files
analyze_pptx('CrystalParadigm_SCIDIP_PitchDeck(by claude).pptx', 'BY_CLAUDE')
analyze_pptx('CrystalParadigm_SCIDIP_PitchDeck_Final.pptx', 'FINAL')
