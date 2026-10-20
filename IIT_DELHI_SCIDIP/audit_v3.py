"""Quick audit of the v3 PPTX — check each slide for shape count, text coverage, and any obvious issues."""
from pptx import Presentation
import sys, codecs
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')

prs = Presentation("CrystalParadigm_SCIDIP_v3_FINAL.pptx")
print(f"Slides: {len(prs.slides)}  |  Dimensions: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
print(f"{'='*90}")

for i, slide in enumerate(prs.slides):
    shapes = list(slide.shapes)
    texts = [s.text.strip() for s in shapes if hasattr(s, 'text') and s.text.strip()]
    total_chars = sum(len(t) for t in texts)
    
    # Check background
    try:
        bg_col = f"#{slide.background.fill.fore_color.rgb}"
    except:
        bg_col = "default"
    
    # Find first meaningful title text
    title_text = ""
    for s in shapes:
        if hasattr(s, 'text') and s.text.strip() and len(s.text.strip()) > 10:
            title_text = s.text.strip()[:70]
            break
    
    print(f"Slide {i+1:2d}  |  Shapes: {len(shapes):3d}  |  Text blocks: {len(texts):2d}  |  "
          f"Chars: {total_chars:5d}  |  BG: {bg_col:10s}  |  '{title_text}'")

print(f"{'='*90}")
print("AUDIT COMPLETE")
