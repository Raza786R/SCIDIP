# -*- coding: utf-8 -*-
from pptx import Presentation
import os

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
PPTX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_PitchDeck.pptx")

prs = Presentation(PPTX_PATH)
fixes = 0

def process_shape(shape):
    global fixes
    if getattr(shape, 'shape_type', None) == 6:  # Group
        for s in shape.shapes:
            process_shape(s)
    
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            original = p.text
            new_text = original
            
            # SLIDE 1
            if "Targeting extreme-environment aerospace propulsion:" in new_text:
                new_text = new_text.replace("Targeting extreme-environment aerospace propulsion:", "Targeting extreme-environment industrial applications:")
            
            # SLIDE 6 & 7
            if "manufacturing_recipe.csv" in new_text:
                new_text = new_text.replace("manufacturing_recipe.csv", "proprietary manufacturing blueprint")
            if "LENS  printing matrix" in new_text:
                pass # it's fine
            
            # SLIDE 9
            if "Accelerated Co-Validation Sprint  & Milestone Framework" in new_text:
                new_text = new_text.replace("Accelerated Co-Validation Sprint  & Milestone Framework", "Validation Plan to TRL 4+ — Incubation Milestone Framework")
            if "MONTHS 1–2" in new_text:
                new_text = new_text.replace("MONTHS 1–2", "COMPUTATIONAL — IN PROGRESS")
            if "MONTHS 3–4" in new_text:
                new_text = new_text.replace("MONTHS 3–4", "REQUIRES LAB ACCESS")
            if "MONTHS 5–6" in new_text:
                new_text = new_text.replace("MONTHS 5–6", "REQUIRES BURNER RIG")
                
            # SLIDE 13
            if "Co-Development Project Structure and Shared Success Criteria" in new_text:
                new_text = "What We Bring and What We Need from SCIDIP"
            if "Partnering through a milestone-funded, shared-risk Joint Development Agreement (JDA) with Safran's hot-section teams." in new_text:
                new_text = "Incubation ask: Access to IIT Delhi's advanced characterization labs and Sona Comstar's industry mentorship."
            if "Crystal Paradigm Commitment" in new_text:
                new_text = "Crystal Paradigm Delivers"
            if "Safran Strategic Support" in new_text:
                new_text = "What We Need from SCIDIP"
            
            if "Execute on-premises multi-physics models" in new_text:
                new_text = "Deliver production-ready AI models, complete multi-physics validation, and provide all digital twin data."
            if "Provide real engine boundary fields and fund the pre-scale sprint" in new_text:
                new_text = "Access to IIT Delhi HPC resources and faculty mentorship on materials characterization methodology."
            if "Scale manufacturing recipes up to complex, subscale high-pressure blade geometries and support joint qualification." in new_text:
                new_text = "Provide 11-layer FGM manufacturing recipes and oversee coupon fabrication & validation."
            if "Provide access to internal testing rigs, engine test cells, and lead downstream flight-readiness certification." in new_text:
                new_text = "Access to high-temperature burner rigs, LENS/DED labs, SEM/EDS/XRD equipment, and TGA facilities."

            if new_text != original:
                # Replace perfectly by overwriting first run and clearing the rest
                if p.runs:
                    p.runs[0].text = new_text
                    for r in p.runs[1:]:
                        r.text = ""
                    fixes += 1
                    print(f"Fixed: {original[:50]} -> {new_text[:50]}")

for slide in prs.slides:
    for shape in slide.shapes:
        process_shape(shape)

prs.save(PPTX_PATH)
print(f"Total precise PPTX fixes applied: {fixes}")
