import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP\CrystalParadigm_SCIDIP_TechnicalAppendix_Final.docx')

print('Looking for section 12...')
for i, p in enumerate(doc.paragraphs):
    if '12. Research Leadership' in p.text:
        print(f'Found at paragraph {i}')
        # Check a few paragraphs before it
        for j in range(max(0, i-10), i):
            print(f'  Para {j}: text="{doc.paragraphs[j].text}"')
            for k, r in enumerate(doc.paragraphs[j].runs):
                if 'lastRenderedPageBreak' in r._element.xml or 'w:br w:type="page"' in r._element.xml:
                    print(f'    Run {k} has a PAGE BREAK in XML')
