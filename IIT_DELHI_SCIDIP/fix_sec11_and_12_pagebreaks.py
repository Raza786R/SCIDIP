import os
from docx import Document
from docx.enum.text import WD_BREAK

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
DOCX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_TechnicalAppendix.docx")

doc = Document(DOCX_PATH)

sec11_idx = -1
sec12_idx = -1

for i, p in enumerate(doc.paragraphs):
    if "11." in p.text and "Computational" in p.text:
        sec11_idx = i
        p.paragraph_format.page_break_before = False
        for run in p.runs:
            if "page" in run._element.xml and "br" in run._element.xml:
                run.clear()
    
    if "12." in p.text and len(p.text) < 40:
        sec12_idx = i
        p.paragraph_format.page_break_before = True

if sec11_idx != -1:
    for i in range(sec11_idx - 1, -1, -1):
        p = doc.paragraphs[i]
        if not p.text.strip():
            p.clear()
            p._element.getparent().remove(p._element)
        else:
            p.paragraph_format.space_after = 12
            break

doc.save(DOCX_PATH)
print("Section 11 pulled up, Section 12 pushed to fresh page.")
