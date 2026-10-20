import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation

BASE_DIR = r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\IIT_DELHI_SCIDIP"
DOCX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_TechnicalAppendix.docx")
PPTX_PATH = os.path.join(BASE_DIR, "CrystalParadigm_SCIDIP_PitchDeck.pptx")
LOGO_PATH = os.path.join(BASE_DIR, "logo.png")

# 1. Extract Logo from PPTX Slide 1
prs = Presentation(PPTX_PATH)
slide1 = prs.slides[0]
logo_extracted = False
for shape in slide1.shapes:
    if shape.shape_type == 13: # msoPicture
        with open(LOGO_PATH, "wb") as f:
            f.write(shape.image.blob)
        logo_extracted = True
        print("Logo extracted from PPTX.")
        break
    if getattr(shape, 'shape_type', None) == 6: # Grouped shape
        for s in shape.shapes:
            if s.shape_type == 13:
                with open(LOGO_PATH, "wb") as f:
                    f.write(s.image.blob)
                logo_extracted = True
                print("Logo extracted from PPTX (Grouped).")
                break
        if logo_extracted:
            break

# 2. Modify DOCX
doc = Document(DOCX_PATH)

# Fix double "advances" and bold "advanced computer vision systems"
for p in doc.paragraphs:
    if "advances inverse design frameworks" in p.text:
        # We need to replace in runs to preserve existing bolding
        for run in p.runs:
            if "advances inverse design" in run.text:
                run.text = run.text.replace("advances inverse design", "developing inverse design")

    if "computer vision systems" in p.text:
        # We want to bold "advanced computer vision systems"
        # Best way is to clear and rebuild, but we might lose other bolds? 
        # Actually this paragraph already has other bolds. Let's do a safe string replacement in runs.
        # Usually it's in a single run if we just added it.
        # We'll just do a targeted run split.
        pass # Will do this more carefully below

def bold_specific_text_in_paragraph(p, search_text):
    # This might mess up if text spans multiple runs, so we re-merge text then re-bold all known highlights.
    pass

# Let's rebuild Raza's paragraph completely to ensure perfect bolding
for p in doc.paragraphs:
    if "Mohammed Ahmed Raza is a Computer Science" in p.text:
        # We know the exact text we want
        full_text = "Mohammed Ahmed Raza is a Computer Science and Engineering student at Neil Gogte Institute of Technology (NGIT) and a Certified AI/ML Professional from iHub-Data, IIIT Hyderabad. He previously served as a Winter Research Intern at IIIT Hyderabad, where he engineered advanced computer vision systems for UAV spatial awareness and obstacle detection. Currently, he advances materials science and deep learning innovations as an AI Research Intern at Crystal Paradigm, where he developing inverse design frameworks for HE-MCGA, computationally engineering next-generation superalloys for extreme aerospace applications. Most recently, he co-authored research introducing a breakthrough Attention Graph Siamese Transformer (DPT-GNN) that drastically optimizes lithium-ion battery capacity predictions through transition-aware learning."
        
        # fix grammar: "where he developing" -> "where he is developing" OR "where he develops"
        # User said "change the second part to 'developing'". So: "where he is developing" or just "developing"
        # "Currently, he advances materials science and deep learning innovations as an AI Research Intern at Crystal Paradigm, developing inverse design frameworks..."
        full_text = full_text.replace("where he developing", "developing")

        p.clear()
        
        highlights = [
            "Certified AI/ML Professional from iHub-Data, IIIT Hyderabad",
            "Winter Research Intern at IIIT Hyderabad",
            "advanced computer vision systems",
            "developing inverse design frameworks for HE-MCGA",
            "Attention Graph Siamese Transformer (DPT-GNN)"
        ]
        
        current_text = full_text
        for h in highlights:
            if h in current_text:
                parts = current_text.split(h)
                p.add_run(parts[0])
                r = p.add_run(h)
                r.bold = True
                if len(parts) > 1:
                    current_text = parts[1]
        if current_text:
            p.add_run(current_text)

# Fix cover page (Center "Crystal Paradigm" and add logo)
for i, p in enumerate(doc.paragraphs):
    if "Crystal Paradigm" in p.text and i < 5: # Assuming it's in the first 5 paragraphs
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Add logo before this paragraph
        if logo_extracted:
            p_logo = p.insert_paragraph_before()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_logo = p_logo.add_run()
            run_logo.add_picture(LOGO_PATH, width=Inches(2.5))
            print("Logo inserted into DOCX.")
        
        # Also center the next few paragraphs on the cover page
        for j in range(i+1, min(i+5, len(doc.paragraphs))):
            doc.paragraphs[j].alignment = WD_ALIGN_PARAGRAPH.CENTER
        break

doc.save(DOCX_PATH)
print("DOCX updated with fixes and logo.")
