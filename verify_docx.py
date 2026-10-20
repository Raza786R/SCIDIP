# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

doc = Document(r"c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\references_and_guides\FINAL_SUBMISSION_FILES\Crystal_Paradigm_Final_Submission.docx")
full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

para_count = len(doc.paragraphs)
table_count = len(doc.tables)
total_words = sum(len(p.text.split()) for p in doc.paragraphs if p.text.strip())

print(f"Paragraphs: {para_count} | Tables: {table_count} | Words: {total_words}")

# IP AUDIT
print("\n=== IP PROTECTION AUDIT ===")
danger = {
    "Nb0.014": "Exact Nb fraction", "Ta0.140": "Exact Ta fraction", "Cr0.236": "Exact Cr fraction",
    "Ta0.16": "Exact Ta (CandA)", "Cr0.27": "Exact Cr (CandA)", "Hf0.18": "Exact Hf (CandA)",
    "github.com": "GitHub link", "fetch_datasets.py": "Script name", "manufacturing_recipe.csv": "Recipe file",
    "Bohrium": "Platform name (should be removed)", "bohrium": "Platform name lowercase",
    "SciMaster": "Platform name",
}
all_safe = True
for ds, desc in danger.items():
    if ds.lower() in full_text.lower():
        print(f"  WARNING: '{ds}' FOUND ({desc})")
        all_safe = False
    else:
        print(f"  SAFE: '{ds}' not found")

# CORRECT VALUES
print("\n=== CORRECT VALUES AUDIT ===")
good = {"1.71": "SF (top candidate)", "7.5 mg/cm": "Oxidation", "0.977": "R2",
        "1400 C": "Temperature", "15.95 MPa": "Shear stress", "1.5981": "1D SF",
        "798 MPa": "Von Mises", "5.02": "VEC", "2727": "Solidus",
        "pending patent": "Patent language", "Challenge #1": "Safran alignment",
        "Challenge #2": "Safran alignment", "X1": "Anonymized composition",
        "Y1": "Anonymized composition", "MACE": "ML Force Field",
        "CHGNet": "ML Force Field", "FEniCS": "FEA solver",
        "CALPHAD": "Phase thermo", "TCHEA4": "Thermo database",
        "9,500": "Budget EUR"}
for gs, desc in good.items():
    if gs in full_text:
        print(f"  PRESENT: '{gs}' ({desc})")
    else:
        print(f"  MISSING: '{gs}' ({desc})")

# Check if 1.66 is also present (for Candidate A)
if "1.66" in full_text:
    print(f"  PRESENT: '1.66' (Candidate A SF)")

if all_safe:
    print("\n=== ALL IP CHECKS PASSED ===")
else:
    print("\n=== WARNING: SOME IP ISSUES FOUND ===")
