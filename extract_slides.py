import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'c:\Users\MOHAMMED AHMED RAZA\Desktop\STUDY\SAFRAN_CHALLENGE_2026\references_and_guides\FINAL_SUBMISSION_FILES\CRYSTAL PARADIGM_ FOR HOTTER, LONGER-LIVED PROPULSION MATERIALS.pptx')

with open('pptx_slides_detailed.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"\n--- SLIDE {i+1} ---\n")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                f.write(shape.text + "\n")
